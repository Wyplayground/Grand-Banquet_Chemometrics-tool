from __future__ import annotations

import html
import json
import math
import os
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Dict, Optional, Tuple

import pandas as pd

from .contracts import (
    CLASSIFICATION_METRICS,
    OVERFIT_DIAGNOSTIC_COLUMNS,
    REGRESSION_METRICS,
    TaskType,
    leaderboard_front_columns,
)


DEEPSEEK_BASE_URL = "https://api.deepseek.com"
DEEPSEEK_MODELS = ("deepseek-flash", "deepseek-v4-pro")
DEEPSEEK_SYSTEM_PROMPT = "You write evidence-grounded chemometrics reports without overstating validation."


def _json_safe(value):
    """Normalize stored evidence without inventing values or emitting NaN/Infinity."""
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (str, int, bool)):
        return value
    if hasattr(value, "tolist"):
        return _json_safe(value.tolist())
    if hasattr(value, "item"):
        return _json_safe(value.item())
    if hasattr(value, "model_dump"):
        return _json_safe(value.model_dump())
    if hasattr(value, "__dict__"):
        return _json_safe(vars(value))
    return str(value)


def structured_summary(
    results: pd.DataFrame,
    task: TaskType,
    top_n: int = 10,
    *,
    context: Optional[Dict[str, object]] = None,
) -> Dict[str, object]:
    """Supply actual metrics for an explicitly limited excerpt of the result table."""
    if isinstance(top_n, bool) or not isinstance(top_n, int) or top_n < 0:
        raise ValueError("top_n must be a non-negative integer.")
    task = TaskType(task)
    eligibility = results.get("Eligible_For_Ranking", pd.Series(False, index=results.index))
    eligible = results[eligibility.fillna(False).eq(True)]
    if "Selection_Rank" in eligible:
        eligible = eligible.sort_values("Selection_Rank", kind="mergesort", na_position="last")
    eligible_count = len(eligible)
    eligible = eligible.head(top_n)
    columns = [
        "Selection_Rank",
        "Pipeline_ID",
        "Preprocessing",
        "Feature_Selection",
        "Model",
        "Selection_Score",
        "Stability_Percentile",
        "Val_Evaluated",
        "Val_Status",
        "Val_Error",
        "CV_Status",
        "CV_Error",
        "Status",
        "Error_Message",
        "Actual_Backend",
        "Hyperparameters",
        "Seed",
        "N_Calibration",
        "N_Validation",
        "Input_Features",
        "Selected_Features",
        "Fit_Seconds",
        "Fold_Primary_SD",
    ]
    metrics = CLASSIFICATION_METRICS if task is TaskType.CLASSIFICATION else REGRESSION_METRICS
    columns.extend(f"{stage}_{metric}" for stage in ("Cal", "CV", "Val") for metric in metrics)
    columns.extend(OVERFIT_DIAGNOSTIC_COLUMNS)
    status = results.get("Status")
    val_status = results.get("Val_Status")
    holdout_failures = None
    if val_status is not None:
        holdout_failures = int(val_status.fillna("").astype(str).str.startswith("FAILED").sum())
    summary = {
        "task": task.value,
        "selection_role": "default ordering uses per-candidate calibration CV evidence (CV_*), not an unbiased evaluation of whole-library selection",
        "validation_role": "Val_* describes internal holdout evaluation; these metric values do not enter the default score",
        "calibration_role": "Cal_* metrics are full-calibration apparent scores and are not used for ranking",
        "default_score_role": "weights and the declared method library are implementation assumptions, not a universal utility function or proof of optimality",
        "total_pipeline_rows": len(results),
        "eligible_pipeline_rows": eligible_count,
        "included_pipeline_rows": len(eligible),
        "requested_top_n": top_n,
        "successful_pipelines": None if status is None else int(status.eq("SUCCESS").sum()),
        "failed_pipelines": None if status is None else int(status.fillna("").astype(str).str.startswith("FAILED").sum()),
        "pipeline_status_counts": None if status is None else status.dropna().astype(str).value_counts().to_dict(),
        "holdout_failed_pipelines": holdout_failures,
        "holdout_status_counts": None if val_status is None else val_status.dropna().astype(str).value_counts().to_dict(),
        "top_pipelines": eligible.reindex(columns=columns).to_dict(orient="records"),
        "context": context,
    }
    return _json_safe(summary)


def build_deepseek_prompt(summary: Dict[str, object]) -> str:
    task = summary["task"]
    metric_guidance = (
        "Interpret supplied Kappa, MCC, macro-F1, balanced accuracy, sensitivity, and specificity values."
        if task == "classification"
        else "Interpret supplied R2, RMSE, MAE, RPD, RPIQ, RER, bias, slope, and intercept values."
    )
    return (
        "You are preparing a technical Grand Banquet chemometrics report. "
        f"The task is {task}. {metric_guidance} "
        "Use only numbers and facts supplied in the JSON below; null means unavailable. "
        "Do not infer missing metrics, group counts, confidence intervals, model versions or superiority. "
        "Do not assume the input includes all 990 configurations; report its included/total row counts. "
        "Clearly distinguish per-candidate calibration CV metrics from internal holdout metrics (Val_*). "
        "Candidate-level CV does not independently assess the whole 990-candidate selection procedure; "
        "highlighting entries after library comparison can introduce selection optimism. "
        "Cal metrics are apparent full-calibration scores only. "
        "Do not use validation metrics to compare or re-rank candidate pipelines; ranking uses CV only. "
        "Default scores and their weights are implementation assumptions; do not call the first entry uniquely best. "
        "Analysts still define the library, tuning, grouping and metrics; never claim analyst bias is eliminated. "
        "Discuss CV performance, fold stability, Cal-vs-CV overfit diagnostics, failed pipelines, "
        "actual model backends, and limitations. "
        "Return a concise academic report with sections: Methods, Ranking Results, Stability and Overfitting, "
        "Internal Holdout Validation, Limitations, and Recommended Next Steps.\n\n"
        + json.dumps(_json_safe(summary), indent=2, ensure_ascii=False, allow_nan=False)
    )


def _redact_record(value, api_key: str):
    """Exclude credentials even when a service error or response echoes the key."""
    if isinstance(value, dict):
        sensitive = {"apikey", "authorization", "access_token", "accesstoken", "password", "secret"}
        return {
            name.replace(api_key, "[REDACTED]") if api_key else name:
            "[REDACTED]" if name.lower().replace("_", "").replace("-", "") in sensitive
            else _redact_record(item, api_key)
            for name, item in value.items()
        }
    if isinstance(value, list):
        return [_redact_record(item, api_key) for item in value]
    if isinstance(value, str) and api_key:
        return value.replace(api_key, "[REDACTED]")
    return value


def _write_call_record(path: Path, record, api_key: str):
    path.write_text(
        json.dumps(_redact_record(_json_safe(record), api_key), ensure_ascii=False, indent=2, allow_nan=False),
        encoding="utf-8",
    )


def call_deepseek(
    summary: Dict[str, object],
    api_key: Optional[str] = None,
    model: str = "deepseek-flash",
    client_factory: Optional[Callable[..., object]] = None,
    *,
    record_dir: Optional[str | Path] = None,
    top_p: Optional[float] = None,
    max_tokens: int = 4096,
    seed: Optional[int] = None,
) -> str:
    """Use DeepSeek only, once, in non-thinking mode; archive evidence without credentials.

    The SDK retries are disabled so a single report action does not trigger
    unrecorded additional attempts. Unsupported seed/top_p arguments are rejected.
    """
    key = (api_key or os.getenv("DEEPSEEK_API_KEY", "")).strip()
    supplied_summary = _json_safe(summary)
    prompt = build_deepseek_prompt(supplied_summary)
    parameters = {"model": model, "temperature": 0.2, "timeout": 120,
                  "max_tokens": max_tokens, "stream": False,
                  "extra_body": {"thinking": {"type": "disabled"}}}
    started = datetime.now(timezone.utc)
    timer = time.perf_counter()
    record = {
        "record_id": uuid.uuid4().hex,
        "status": "STARTED",
        "started_at_utc": started.isoformat(),
        "provider": "DeepSeek",
        "endpoint": DEEPSEEK_BASE_URL,
        "client_parameters": {"base_url": DEEPSEEK_BASE_URL, "timeout": 120, "max_retries": 0},
        "credential_policy": "Credentials are excluded; accidental credential echoes in inputs, errors or responses are redacted.",
        "summary": supplied_summary,
        "system_prompt": DEEPSEEK_SYSTEM_PROMPT,
        "user_prompt": prompt,
        "requested_parameters": parameters,
        "unsupported_parameter_inputs": {"top_p": top_p, "seed": seed},
        "api_request_attempts": 0,
        "response": None,
        "error": None,
    }
    record_path = None
    if record_dir is not None:
        directory = Path(record_dir)
        directory.mkdir(parents=True, exist_ok=True)
        record_path = directory / f"deepseek_call_{started:%Y%m%dT%H%M%S_%fZ}_{record['record_id']}.json"
        _write_call_record(record_path, record, key)
    try:
        if model not in DEEPSEEK_MODELS:
            raise ValueError("Unsupported DeepSeek model; select deepseek-flash or deepseek-v4-pro.")
        if seed is not None:
            raise ValueError("DeepSeek report generation does not support a seed parameter.")
        if top_p is not None:
            raise ValueError("top_p is not sent in DeepSeek non-thinking report generation.")
        if isinstance(max_tokens, bool) or not isinstance(max_tokens, int) or not 1 <= max_tokens <= 393216:
            raise ValueError("max_tokens must be an integer between 1 and 393216.")
        if not key:
            raise ValueError("A DeepSeek API key is required through the GUI or DEEPSEEK_API_KEY.")
        if client_factory is None:
            try:
                from openai import OpenAI
            except ImportError as exc:
                raise ImportError("The openai package is required for DeepSeek report generation.") from exc
            client_factory = OpenAI
        client = client_factory(api_key=key, base_url=DEEPSEEK_BASE_URL, timeout=120, max_retries=0)
        record["api_request_attempts"] = 1
        completion = client.chat.completions.create(
            messages=[
                {"role": "system", "content": DEEPSEEK_SYSTEM_PROMPT},
                {"role": "user", "content": prompt},
            ],
            **parameters,
        )
        raw_response = _json_safe(completion)
        record["response"] = {
            "returned_model": getattr(completion, "model", None),
            "response_id": getattr(completion, "id", None),
            "request_id": getattr(completion, "_request_id", None),
            "created": getattr(completion, "created", None),
            "usage": _json_safe(getattr(completion, "usage", None)),
            "text": None,
            "finish_reason": None,
            "raw_payload": raw_response,
        }
        choices = getattr(completion, "choices", None)
        if not choices:
            raise RuntimeError("DeepSeek returned no report choice.")
        choice = choices[0]
        text = getattr(getattr(choice, "message", None), "content", None)
        finish_reason = getattr(choice, "finish_reason", None)
        record["response"]["text"] = text
        record["response"]["finish_reason"] = finish_reason
        if not text or not str(text).strip():
            raise RuntimeError("DeepSeek returned an empty report.")
        if finish_reason != "stop":
            raise RuntimeError(f"DeepSeek report was not completed normally (finish_reason={finish_reason!r}).")
        record["status"] = "SUCCESS"
        return str(text).strip()
    except Exception as exc:
        record["status"] = "FAILED"
        record["error"] = {
            "type": type(exc).__name__,
            "message": str(exc),
            "status_code": getattr(exc, "status_code", None),
            "request_id": getattr(exc, "request_id", None),
        }
        raise
    finally:
        record["completed_at_utc"] = datetime.now(timezone.utc).isoformat()
        record["elapsed_seconds"] = time.perf_counter() - timer
        if record_path is not None:
            _write_call_record(record_path, record, key)


def write_report_files(text: str, output_dir: str | Path, stem: str = "gb_ai_report") -> Tuple[Path, Path, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ") + "_" + uuid.uuid4().hex
    text_path = output_dir / f"{stem}_{timestamp}.md"
    docx_path = output_dir / f"{stem}_{timestamp}.docx"
    pdf_path = output_dir / f"{stem}_{timestamp}.pdf"
    text_path.write_text(text, encoding="utf-8")

    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("python-docx is required for Word report export.") from exc
    document = Document()
    document.add_heading("Grand Banquet Chemometrics Report", level=0)
    for block in text.split("\n\n"):
        stripped = block.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            document.add_heading(stripped.lstrip("# "), level=min(3, len(stripped) - len(stripped.lstrip("#"))))
        else:
            document.add_paragraph(stripped)
    document.save(docx_path)

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise ImportError("reportlab is required for PDF report export.") from exc
    styles = getSampleStyleSheet()
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    for style_name in ("Title", "BodyText"):
        styles[style_name].fontName = "STSong-Light"
    story = [Paragraph("Grand Banquet Chemometrics Report", styles["Title"]), Spacer(1, 6 * mm)]
    for block in text.split("\n\n"):
        if block.strip():
            safe = block.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
            story.extend([Paragraph(safe, styles["BodyText"]), Spacer(1, 3 * mm)])
    SimpleDocTemplate(str(pdf_path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm).build(story)
    return text_path, docx_path, pdf_path


def write_leaderboard_html(results: pd.DataFrame, destination: str | Path, task: TaskType) -> Path:
    destination = Path(destination)
    task = TaskType(task)
    front = [column for column in leaderboard_front_columns(task) if column in results.columns]
    remainder = [column for column in results.columns if column not in front]
    ordered = results.loc[:, front + remainder].copy()
    # Links to CV prediction files when Pipeline_ID is present.
    if "Pipeline_ID" in ordered.columns:
        ordered.insert(
            0,
            "CV_Predictions",
            ordered["Pipeline_ID"].map(
                lambda pipeline_id: f"<a href='predictions/{html.escape(str(pipeline_id))}_cv.csv'>cv</a>"
            ),
        )
    table = ordered.to_html(index=False, border=0, escape=False, na_rep="")
    preprocessings = sorted({str(value) for value in results.get("Preprocessing", pd.Series(dtype=object)).dropna()})
    selectors = sorted({str(value) for value in results.get("Feature_Selection", pd.Series(dtype=object)).dropna()})
    models = sorted({str(value) for value in results.get("Model", pd.Series(dtype=object)).dropna()})
    statuses = sorted({str(value) for value in results.get("Status", pd.Series(dtype=object)).dropna()})

    def options(values):
        return "".join(f"<option value='{html.escape(value)}'>{html.escape(value)}</option>" for value in values)

    content = f"""<!doctype html>
<html>
<head>
<meta charset='utf-8'>
<title>Grand Banquet v3 Leaderboard</title>
<style>
body{{font-family:Segoe UI,sans-serif;margin:24px;background:#E0F7FA;color:#263238}}
h1{{color:#00796B}}
.filters{{display:flex;flex-wrap:wrap;gap:12px;margin:12px 0 18px;padding:12px;background:#FFFFFF;border:1px solid #B0BEC5}}
.filters label{{font-size:12px;display:flex;flex-direction:column;gap:4px}}
.table-wrap{{overflow:auto;max-height:78vh;background:white;border:1px solid #B0BEC5}}
table{{border-collapse:collapse;min-width:100%;background:white}}
th{{background:#00796B;color:white;position:sticky;top:0;z-index:1}}
th,td{{border:1px solid #B0BEC5;padding:6px;font-size:12px;white-space:nowrap}}
tr.rank1{{background:#FFF8E1}}
tr.hidden{{display:none}}
.note{{font-size:13px;max-width:960px}}
</style>
</head>
<body>
<h1>Grand Banquet Selection Leaderboard</h1>
<p class='note'>The default score orders candidates using per-candidate calibration CV evidence (CV_*) and stated implementation weights. It is not an unbiased assessment of whole-library selection or proof of a uniquely best pipeline. Cal_* scores are full-calibration apparent metrics. Internal holdout metric values (Val_*) do not enter the default score. Overfit_* columns are descriptive diagnostics only; analysts still define the library, tuning, grouping and metrics.</p>
<div class='filters'>
<label>Preprocessing<select id='filter-pre'><option value=''>All</option>{options(preprocessings)}</select></label>
<label>Feature selection<select id='filter-sel'><option value=''>All</option>{options(selectors)}</select></label>
<label>Model<select id='filter-model'><option value=''>All</option>{options(models)}</select></label>
<label>Status<select id='filter-status'><option value=''>All</option>{options(statuses)}</select></label>
</div>
<div class='table-wrap'>{table}</div>
<script>
const col = (name) => {{
  const headers = Array.from(document.querySelectorAll('table thead th')).map(th => th.textContent.trim());
  return headers.indexOf(name);
}};
const preCol = col('Preprocessing');
const selCol = col('Feature_Selection');
const modelCol = col('Model');
const statusCol = col('Status');
const rankCol = col('Selection_Rank');
document.querySelectorAll('table tbody tr').forEach(row => {{
  if (rankCol >= 0 && row.children[rankCol] && row.children[rankCol].textContent.trim() === '1') {{
    row.classList.add('rank1');
  }}
}});
function applyFilters() {{
  const pre = document.getElementById('filter-pre').value;
  const sel = document.getElementById('filter-sel').value;
  const model = document.getElementById('filter-model').value;
  const status = document.getElementById('filter-status').value;
  document.querySelectorAll('table tbody tr').forEach(row => {{
    const cells = row.children;
    const ok =
      (!pre || (preCol >= 0 && cells[preCol].textContent.trim() === pre)) &&
      (!sel || (selCol >= 0 && cells[selCol].textContent.trim() === sel)) &&
      (!model || (modelCol >= 0 && cells[modelCol].textContent.trim() === model)) &&
      (!status || (statusCol >= 0 && cells[statusCol].textContent.trim() === status));
    row.classList.toggle('hidden', !ok);
  }});
}}
['filter-pre','filter-sel','filter-model','filter-status'].forEach(id => {{
  document.getElementById(id).addEventListener('change', applyFilters);
}});
</script>
</body>
</html>
"""
    destination.write_text(content, encoding="utf-8")
    return destination
